"""Real subprocess export, timeout/cancel, and selection-to-open UI integration."""
import os
os.environ.setdefault('QT_QPA_PLATFORM','offscreen')
import sys
import tempfile
from pathlib import Path
from unittest.mock import patch
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from PIL import Image
from pptx import Presentation
from PySide6.QtCore import QTimer, QProcess
from PySide6.QtWidgets import QApplication, QDialog, QMessageBox
from PySide6.QtGui import QDesktopServices
from ppt_export_dialog import PptExportDialog
from catalog import Catalog
from app import Window, Service

app=QApplication([])
with tempfile.TemporaryDirectory() as temp:
    root=Path(temp)
    sources=root/'sources';sources.mkdir()
    image=sources/'source.tif'
    pages=[Image.new('RGB',(720,480),color) for color in ('red','green','blue')]
    pages[0].save(image,save_all=True,append_images=pages[1:])
    for page in pages:page.close()
    original=image.read_bytes()
    dialog=PptExportDialog([image],15,300)
    assert dialog.exec()==QDialog.DialogCode.Accepted, dialog.error
    assert dialog.result_info['pages']==3
    assert len(Presentation(dialog.output).slides[0].shapes)==3
    assert dialog.output.exists() and not list(dialog.directory.glob('images-*'))
    assert dialog.process.state()==QProcess.ProcessState.NotRunning
    import shutil
    shutil.rmtree(dialog.directory)

    slow=root/'slow.py'
    slow.write_text('import time\ntime.sleep(30)\n')
    for mode in ('timeout','cancel','close'):
        dialog=PptExportDialog([image],.2 if mode=='timeout' else 15,300,worker_script=slow)
        ticks=[];timer=QTimer();timer.timeout.connect(lambda:ticks.append(1));timer.start(20)
        if mode=='cancel':QTimer.singleShot(200,dialog.reject)
        if mode=='close':QTimer.singleShot(200,dialog.close)
        assert dialog.exec()==QDialog.DialogCode.Rejected
        timer.stop()
        assert len(ticks)>=3, 'GUI event loop blocked'
        assert dialog.process.state()==QProcess.ProcessState.NotRunning
        assert not dialog.directory.exists()
        assert bool(dialog.error)==(mode=='timeout')
    dialog=PptExportDialog([root/'missing.tif'],15,300)
    assert dialog.exec()==QDialog.DialogCode.Rejected
    assert 'FileNotFoundError' in dialog.error

    db=Catalog(root/'db'/'catalog.sqlite3');db.add_root(sources);db.scan(db.roots()[0])
    with patch.object(Service,'start'):window=Window(db)
    try:
        captured=[]
        with patch.object(QDesktopServices,'openUrl',side_effect=lambda url:captured.append(url.toLocalFile()) or True),patch.object(QMessageBox,'information'):
            window.export_images_to_ppt(db.rows())
        assert len(captured)==1
        output=Path(captured[0]);assert len(Presentation(output).slides[0].shapes)==3
        assert window.ppt_export_dialog is None
        assert image.read_bytes()==original
        shutil.rmtree(output.parent)
    finally:
        window.details.timer.stop();window.details.save_timer.stop()
        window.stats_timer.stop();window.refresh_timer.stop();window.db_poll.stop()
        window.ui_pool.shutdown(wait=True);window.model.pool.shutdown(wait=True)
        window.exiting=True;window.tray.hide();window.close()
print('PPT export GUI / timeout / cancel / close smoke passed')
