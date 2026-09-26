from pathlib import Path
import io
import sys
import tempfile
import unittest
import zipfile
import numpy as np
from PIL import Image, TiffImagePlugin
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from ppt_export import export_presentation, layout_images, display_image


class PptExportTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.base=Path(self.temp.name)

    def tiff(self, count=3, orientation=1, ome=None):
        pages=[Image.fromarray((np.arange(120*80).reshape(80,120)+i).astype(np.uint16)) for i in range(count)]
        tags=TiffImagePlugin.ImageFileDirectory_v2()
        tags[282]=20000; tags[283]=10000; tags[296]=3; tags[274]=orientation
        if ome: tags[270]=ome
        path=self.base/'stack.tif'
        pages[0].save(path,save_all=True,append_images=pages[1:],tiffinfo=tags)
        for page in pages: page.close()
        return path

    def test_all_pages_full_pixels_native_scale_and_readonly(self):
        source=self.tiff(23)
        before=source.read_bytes()
        output=self.base/'out.pptx'
        info=export_presentation([source],output)
        self.assertEqual(info['pages'],23)  # exceeds thumbnail cap
        self.assertEqual(info['normalized'],23)
        self.assertEqual(info['uncalibrated'],[])
        self.assertEqual(source.read_bytes(),before)
        deck=Presentation(output)
        self.assertEqual(len(deck.slides),1)
        groups=deck.slides[0].shapes
        self.assertEqual(len(groups),23)
        for group in groups:
            self.assertEqual(group.shape_type,MSO_SHAPE_TYPE.GROUP)
            picture,text,line=group.shapes
            self.assertEqual(picture.image.size,(120,80))
            self.assertEqual(line.shape_type,MSO_SHAPE_TYPE.LINE)
            self.assertEqual(text.text,'10 µm')
            self.assertAlmostEqual(line.width/picture.width,10/(120*.5),places=5)
            self.assertLessEqual(group.left+group.width,deck.slide_width)
            self.assertLessEqual(group.top+group.height,deck.slide_height)
        with zipfile.ZipFile(output) as z:
            self.assertIn(b'autoCompressPictures="0"',z.read('ppt/presentation.xml'))

    def test_single_jpeg_and_png_bytes_preserved_and_mixed_sizes(self):
        jpg=self.base/'photo.jpg'; png=self.base/'tall.png'
        Image.new('RGB',(2400,900),'red').save(jpg)
        Image.new('RGBA',(400,1600),(0,200,0,80)).save(png)
        tiff=self.tiff(1)
        output=self.base/'mixed.pptx'
        info=export_presentation([jpg,png,tiff],output)
        self.assertEqual(info['pages'],3)
        self.assertEqual(len(info['uncalibrated']),2)
        deck=Presentation(output)
        pictures=[g.shapes[0] for g in deck.slides[0].shapes]
        self.assertEqual(pictures[0].image.blob,jpg.read_bytes())
        self.assertEqual(pictures[1].image.blob,png.read_bytes())
        for p in pictures:
            iw,ih=p.image.size
            self.assertAlmostEqual(p.width/p.height,iw/ih,places=5)

    def test_ome_overrides_resolution(self):
        xml='<OME><Image><Pixels SizeX="120" SizeY="80" PhysicalSizeX="2" PhysicalSizeY="3"><TiffData/></Pixels></Image></OME>'
        output=self.base/'ome.pptx'
        export_presentation([self.tiff(2,ome=xml)],output)
        for g in Presentation(output).slides[0].shapes:
            picture,text,line=g.shapes
            self.assertEqual(text.text,'50 µm')
            self.assertAlmostEqual(line.width/picture.width,50/(120*2),places=5)

    def test_orientation_swaps_pixels_and_calibration_axes(self):
        output=self.base/'rotated.pptx'
        export_presentation([self.tiff(1,orientation=6)],output)
        picture,text,line=Presentation(output).slides[0].shapes[0].shapes
        self.assertEqual(picture.image.size,(80,120))
        self.assertEqual(text.text,'20 µm')
        self.assertAlmostEqual(line.width/picture.width,20/(80*1),places=5)

    def test_tiff_orientation_pixels_with_and_without_compression(self):
        raw=np.arange(7*11,dtype=np.uint8).reshape(7,11)
        transforms={2:Image.Transpose.FLIP_LEFT_RIGHT,3:Image.Transpose.ROTATE_180,
                    4:Image.Transpose.FLIP_TOP_BOTTOM,5:Image.Transpose.TRANSPOSE,
                    6:Image.Transpose.ROTATE_270,7:Image.Transpose.TRANSVERSE,
                    8:Image.Transpose.ROTATE_90}
        for compression in ('raw','tiff_lzw'):
            for orientation in range(1,9):
                with self.subTest(compression=compression,orientation=orientation):
                    source=Image.fromarray(raw)
                    path=self.base/'orientation.tif'
                    source.save(path,compression=compression,tiffinfo={274:orientation})
                    expected=source.transpose(transforms[orientation]) if orientation in transforms else source.copy()
                    with Image.open(path) as im:
                        result,_=display_image(im,orientation)
                        np.testing.assert_array_equal(np.asarray(result),np.asarray(expected))
                        result.close()
                    source.close();expected.close()

    def test_no_overwrite_pixel_limit_and_invalid_source(self):
        existing=self.base/'out.pptx'; existing.write_bytes(b'keep')
        with self.assertRaises(FileExistsError): export_presentation([self.tiff()],existing)
        self.assertEqual(existing.read_bytes(),b'keep')
        big=self.base/'big.png'; Image.new('L',(1100,1100)).save(big)
        with self.assertRaises(Image.DecompressionBombError): export_presentation([big],self.base/'big.pptx',1)
        self.assertFalse((self.base/'big.pptx').exists())
        with self.assertRaises(FileNotFoundError): export_presentation([self.base/'missing.tif'],self.base/'missing.pptx')
        self.assertFalse(list(self.base.glob('images-*')))

    def test_layout_nonoverlap_for_many_and_extreme_aspects(self):
        for count in (1,3,8,21,101):
            pages=[dict(width=w,height=h) for w,h in ([(100,100),(10,1000),(1000,10)]*count)[:count]]
            rects=layout_images(pages,12000000,7000000)
            for i,(x,y,w,h) in enumerate(rects):
                self.assertGreater(w,0); self.assertGreater(h,0)
                self.assertGreaterEqual(x,0); self.assertGreaterEqual(y,0)
                self.assertLessEqual(x+w,12000000); self.assertLessEqual(y+h,7000000)
                for a,b,c,d in rects[i+1:]:
                    self.assertTrue(x+w<=a or a+c<=x or y+h<=b or b+d<=y)


if __name__=='__main__': unittest.main()
