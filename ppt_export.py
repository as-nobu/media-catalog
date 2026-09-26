"""Explicit full-resolution image export. Runs in an isolated, cancellable process."""
import argparse
import json
import math
from pathlib import Path
import sys
import tempfile
from types import SimpleNamespace

from PIL import Image, ImageOps
from thumbnail_worker import configure_image_limit, copy_source
from tiff_scale import calibration, nice_bar, ome_scales


def display_image(image, tiff_orientation=None):
    """Keep all pixels; normalize scientific grayscale only for PowerPoint display."""
    if tiff_orientation is not None:
        # Pillow versions differ in whether TIFF .size is already oriented before
        # decoding. Decode the raw tile dimensions, then apply orientation once.
        raw_size = (int(image.tag_v2[256]), int(image.tag_v2[257]))
        image.tag_v2[274] = 1
        image.getexif()[274] = 1
        image._size = raw_size
        image.load()
        transforms = {2:Image.Transpose.FLIP_LEFT_RIGHT,3:Image.Transpose.ROTATE_180,
                      4:Image.Transpose.FLIP_TOP_BOTTOM,5:Image.Transpose.TRANSPOSE,
                      6:Image.Transpose.ROTATE_270,7:Image.Transpose.TRANSVERSE,
                      8:Image.Transpose.ROTATE_90}
        page = image.transpose(transforms[tiff_orientation]) if tiff_orientation in transforms else image.copy()
    else:
        image.load()
        page = ImageOps.exif_transpose(image)
    if page.mode.startswith('I') or page.mode == 'F':
        import numpy as np
        values = np.asarray(page, dtype=np.float64)
        finite = np.isfinite(values)
        output = np.zeros(values.shape, dtype=np.uint8)
        if finite.any():
            low, high = values[finite].min(), values[finite].max()
            if high > low:
                output[finite] = np.clip((values[finite]-low)*255/(high-low), 0, 255).astype(np.uint8)
        page.close()
        return Image.fromarray(output), True
    if page.mode not in ('RGB', 'RGBA', 'L', 'LA'):
        converted = page.convert('RGBA' if 'transparency' in page.info else 'RGB')
        page.close()
        page = converted
    return page, False


def prepare_images(paths, directory, progress):
    pages = []
    normalized = 0
    for file_index, path in enumerate(paths, 1):
        path = Path(path)
        progress('source', file=file_index, files=len(paths), name=path.name)
        staged = directory / ('source-'+str(file_index)+path.suffix)
        # Only this explicit export operation opens originals, in read-only mode.
        copy_source(path, staged)
        with Image.open(staged) as image:
            is_tiff = image.format == 'TIFF'
            count = image.n_frames if is_tiff else 1  # no thumbnail page cap
            ome = ome_scales(image.tag_v2.get(270), count) if is_tiff else {}
            for index in range(count):
                image.seek(index)
                # Capture orientation/calibration before Pillow load mutates TIFF tags.
                if is_tiff:
                    raw_size = (int(image.tag_v2[256]), int(image.tag_v2[257]))
                    raw = SimpleNamespace(tag_v2=image.tag_v2,size=raw_size,width=raw_size[0],height=raw_size[1])
                    scale = calibration(raw, ome.get(index))
                    orientation = image.tag_v2.get(274,1)
                else:
                    scale = {}
                    orientation = image.getexif().get(274, 1)
                reusable = image.format in ('JPEG', 'PNG') and orientation == 1 and image.mode in ('RGB','RGBA','L','LA','P')
                if reusable:
                    image.load()  # validate the image even when embedding its original bytes
                    output = staged
                    size = image.size
                else:
                    page, converted = display_image(image, orientation if is_tiff else None)
                    normalized += int(converted)
                    output = directory / f'page-{len(pages)+1}.png'
                    try:
                        # Do not propagate TIFF/EXIF orientation into a transposed PNG.
                        page.info.clear()
                        page.save(output, 'PNG')
                        size = page.size
                    finally:
                        page.close()
                if scale and (scale.get('width'), scale.get('height')) != size:
                    raise ValueError('Image orientation and calibration dimensions disagree: '+path.name)
                pages.append(dict(image=str(output), width=size[0], height=size[1],
                                  scale=scale, name=path.name, page=index+1))
                progress('page', file=file_index, files=len(paths), name=path.name, page=index+1, pages=count)
    return pages, normalized


def layout_images(pages, width, height):
    """Choose a uniform grid using actual aspect ratios, without cropping."""
    if not pages:
        raise ValueError('No images selected.')
    margin = int(min(width, height)*.03)
    available_w, available_h = width-2*margin, height-2*margin
    def candidate(columns):
        rows = math.ceil(len(pages)/columns)
        gap = min(int(min(width,height)*.012), available_w//(columns*10), available_h//(rows*10))
        cell_w = (available_w-gap*(columns-1))/columns
        cell_h = (available_h-gap*(rows-1))/rows
        result = []
        for index, page in enumerate(pages):
            factor = min(cell_w/page['width'], cell_h/page['height'])
            w, h = max(1, int(page['width']*factor)), max(1, int(page['height']*factor))
            x = margin+(index%columns)*(cell_w+gap)+(cell_w-w)/2
            y = margin+(index//columns)*(cell_h+gap)+(cell_h-h)/2
            result.append((int(x),int(y),w,h))
        return result
    return max((candidate(c) for c in range(1,len(pages)+1)),
               key=lambda rects: sum(w*h for _,_,w,h in rects))


def add_scale_bar(group, scale, rect):
    from pptx.enum.shapes import MSO_CONNECTOR
    from pptx.enum.text import PP_ALIGN
    from pptx.dml.color import RGBColor
    from pptx.util import Pt
    x,y,w,h = rect
    bar = nice_bar(scale, w)
    if bar is None:
        return False
    length, label = bar
    boxw = int(w*.40)
    font = min(14, boxw/12700/(len(label)*.7+1), h/12700*.065)
    boxh = max(1,int(Pt(font*2.5)))
    pad = min(int(w*.025),int(h*.025))
    left, top = x+w-boxw-pad, y+h-boxh-pad
    text = group.shapes.add_textbox(left,top,boxw,boxh)
    text.name = 'Scale label: '+label
    text.fill.solid(); text.fill.fore_color.rgb = RGBColor(0,0,0)
    text.line.fill.background()
    frame = text.text_frame
    frame.word_wrap = False
    frame.margin_left = frame.margin_right = frame.margin_bottom = 0
    frame.margin_top = int(Pt(font*.95))
    paragraph = frame.paragraphs[0]
    paragraph.text = label
    paragraph.alignment = PP_ALIGN.CENTER
    paragraph.font.size = Pt(font)
    paragraph.font.name = 'Arial'
    paragraph.font.color.rgb = RGBColor(255,255,255)
    start = int(left+(boxw-length)/2)
    line_y = top+int(Pt(font*.5))
    line = group.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,start,line_y,start+int(length),line_y)
    line.name = 'Scale bar: '+label
    line.line.color.rgb = RGBColor(255,255,255)
    line.line.width = Pt(min(2,font*.16))
    return True


def export_presentation(paths, output, max_image_mp=300, progress=lambda *a, **k: None):
    from pptx import Presentation
    from pptx.util import Inches
    configure_image_limit(max_image_mp)
    output = Path(output)
    if output.exists():
        raise FileExistsError(str(output))
    if not paths:
        raise ValueError('No images selected.')
    with tempfile.TemporaryDirectory(prefix='images-',dir=output.parent) as temp:
        pages, normalized = prepare_images(paths, Path(temp), progress)
        deck = Presentation()
        deck.slide_width, deck.slide_height = Inches(13.333333), Inches(7.5)
        deck._element.set('autoCompressPictures','0')
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        uncalibrated = []
        for page, rect in zip(pages, layout_images(pages,deck.slide_width,deck.slide_height)):
            group = slide.shapes.add_group_shape()
            group.name = f"{page['name']} [{page['page']}]"
            picture = group.shapes.add_picture(page['image'],*rect)
            picture.name = group.name
            if not add_scale_bar(group,page['scale'],rect):
                uncalibrated.append(group.name)
        slide.notes_slide.notes_text_frame.text = '\n'.join(
            f"{p['name']} [{p['page']}]" for p in pages)
        progress('saving', pages=len(pages))
        # Output lives in a unique temporary directory owned by the export dialog.
        # Reserve exclusively: never overwrite an existing file.
        with output.open('xb') as stream:
            deck.save(stream)
    return dict(pages=len(pages), uncalibrated=uncalibrated, normalized=normalized)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('request')
    args = parser.parse_args()
    request = json.loads(Path(args.request).read_text(encoding='utf-8'))
    def progress(stage, **values):
        print(json.dumps(dict(stage=stage,**values),ensure_ascii=True),flush=True)
    try:
        result = export_presentation(request['paths'],request['output'],request['max_image_mp'],progress)
        progress('done',**result)
    except Exception as exc:
        print(type(exc).__name__+': '+str(exc),file=sys.stderr)
        return 1
    return 0


if __name__ == '__main__':
    sys.exit(main())
